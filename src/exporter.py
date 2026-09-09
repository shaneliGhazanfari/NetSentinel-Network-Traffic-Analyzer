import os
import pandas as pd
from pathlib import Path
from openpyxl.styles import Alignment
import loader as loader


def export_analysis(df):
    BASE_DIR = Path(__file__).resolve().parent.parent
    output_dir = BASE_DIR / "output"
    os.makedirs(output_dir, exist_ok=True)

    df_analyze = loader.analyze_dataset(df)
    suspicious_df = loader.detect_suspicious_ips(df)
    risk_df = loader.calculate_risk_score(suspicious_df)

    csv_path = output_dir / "security_analysis.csv"
    risk_df.to_csv(csv_path, index=False)

    excel_path = output_dir / "security_analysis.xlsx"

    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:

        risk_df.to_excel(writer, sheet_name="Risk Analysis", index=False)

        analysis_data = {
            "Metric": [
                "Dataset Rows",
                "Dataset Columns",
                "Duplicate Count",
                "Unique Source IPs",
                "Unique Destination IPs",
            ],
            "Value": [
                df_analyze["dataset_rows"],
                df_analyze["dataset_columns"],
                df_analyze["duplicate_count"],
                df_analyze["unique_source_ips"],
                df_analyze["unique_destination_ips"],
            ],}

        analysis_df = pd.DataFrame(analysis_data)
        analysis_df.to_excel(writer, sheet_name="Dataset Analysis", index=False)

        top_risk_ips = risk_df.sort_values( by="risk_score",ascending=False).head(5)

        workbook = writer.book
        report_sheet = workbook.create_sheet("Security Report")

        report_sheet["A1"] = "NetSentinel Security Report"

        report_sheet["A3"] = "Summary"

        report_sheet["A4"] = "Total Suspicious IPs"
        report_sheet["B4"] = risk_df.shape[0]

        report_sheet["A5"] = "High Risk"
        report_sheet["B5"] = len(risk_df[risk_df["risk_level"] == "High"])

        report_sheet["A6"] = "Medium Risk"
        report_sheet["B6"] = len(risk_df[risk_df["risk_level"] == "Medium"])

        report_sheet["A7"] = "Low Risk"
        report_sheet["B7"] = len(risk_df[risk_df["risk_level"] == "Low"])

        report_sheet["A8"] = "Normal"
        report_sheet["B8"] = len(risk_df[risk_df["risk_level"] == "Normal"])

        report_sheet["A10"] = "Average Risk Score"
        report_sheet["B10"] = round(risk_df["risk_score"].mean(), 2)

        report_sheet["A11"] = "Maximum Risk Score"
        report_sheet["B11"] = risk_df["risk_score"].max()

        report_sheet["A13"] = "High Risk IPs"

        row = 14
        for ip in risk_df[risk_df["risk_level"] == "High"]["srcip"]:
            report_sheet.cell(row=row, column=1).value = ip
            row += 1

        row += 2
        report_sheet.cell(row=row, column=1).value = "Top 5 Highest Risk IPs"

        row += 1
        report_sheet.cell(row=row, column=1).value = "IP Address"
        report_sheet.cell(row=row, column=2).value = "Risk Score"
        report_sheet.cell(row=row, column=3).value = "Risk Level"
        report_sheet.cell(row=row, column=4).value = "Risk Reasons"

        for _, item in top_risk_ips.iterrows():
            row += 1
            report_sheet.cell(row=row, column=1).value = item["srcip"]
            report_sheet.cell(row=row, column=2).value = item["risk_score"]
            report_sheet.cell(row=row, column=3).value = item["risk_level"]
            report_sheet.cell(row=row, column=4).value = ", ".join(item["risk_reasons"])


        risk_sheet = workbook["Risk Analysis"]
        risk_sheet.freeze_panes = "A2"
        risk_sheet.auto_filter.ref = risk_sheet.dimensions

        for column in risk_sheet.columns:
            max_length = 0
            column_letter = column[0].column_letter

            for cell in column:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))

            risk_sheet.column_dimensions[column_letter].width = min(max_length + 2, 40)

        analysis_sheet = workbook["Dataset Analysis"]
        analysis_sheet.freeze_panes = "A2"
        analysis_sheet.column_dimensions["A"].width = 30
        analysis_sheet.column_dimensions["B"].width = 25

        report_sheet.freeze_panes = "A4"

        report_sheet.column_dimensions["A"].width = 30
        report_sheet.column_dimensions["B"].width = 15
        report_sheet.column_dimensions["C"].width = 15
        report_sheet.column_dimensions["D"].width = 45

        for cell in report_sheet["D"]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    print(f"CSV report exported to: {csv_path}")
    print(f"Excel report exported to: {excel_path}")

    return csv_path, excel_path